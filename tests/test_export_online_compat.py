"""Tests for util.export_online_compat."""

from __future__ import annotations

import csv
import json
from pathlib import Path

import pytest

from util.export_online_compat import transform_delivery_csv_to_xlsx

pytest.importorskip("openpyxl")

_REPO = Path(__file__).resolve().parents[1]
_ROI = _REPO / "EXAMPLES" / "RoI"


def test_transform_roi_config_smoke(tmp_path: Path) -> None:
    config_path = _ROI / "export_config.json"
    if not config_path.is_file():
        pytest.skip("RoI export_config.json not present")

    stocking_title = json.loads(config_path.read_text(encoding="utf-8"))["column_sources"]["112"]
    csv_path = tmp_path / "sample.csv"
    with csv_path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["File", "Supplier Name", "Spouse_Partner_Full_time", stocking_title])
        writer.writerow(["PDF/0001.pdf", "Acme Farm", "Ticked", "Stocking Rate: 170-220kg N/ha"])

    template = _ROI / "Planning Census RoI 2026-2031-Final_REDACTED.xlsx"
    if not template.is_file():
        templates = list(_ROI.glob("*.xlsx"))
        if not templates:
            pytest.skip("RoI Qualtrics template .xlsx not present")
        template = templates[0]

    out = transform_delivery_csv_to_xlsx(
        csv_path=csv_path,
        config_path=config_path,
        template_path=template,
    )
    assert out.is_file()
    assert out.suffix == ".xlsx"
    assert out.name == "sample_for_import.xlsx"

    from openpyxl import load_workbook

    wb = load_workbook(out, read_only=True)
    try:
        rows = list(wb.active.iter_rows(values_only=True))
    finally:
        wb.close()
    assert len(rows) >= 2
    assert rows[0][0] == "File"
    assert rows[0][1] == "Respondent ID"
    assert rows[2][0] == "PDF/0001.pdf"

    spouse_col = rows[1].index("Spouse/Partner - Full time")
    assert rows[2][spouse_col] == "Full time"
    stocking_col = next(
        i for i, stem in enumerate(rows[0]) if (stem or "").startswith("What is your Whole Farm Stocking Rate")
    )
    assert rows[2][stocking_col] == "170-220kg N/ha"
