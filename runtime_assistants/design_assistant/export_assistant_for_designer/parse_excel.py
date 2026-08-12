"""Parse Qualtrics-style Excel export headers into an export format descriptor."""

from __future__ import annotations

import logging
from pathlib import Path

from runtime_assistants.design_assistant.export_assistant_for_designer.schema import (
    ExportColumn,
    ExportFormatDescriptor,
    ExportGroup,
    KIND_META,
    KIND_MULTI_SELECT_OPTION,
    KIND_OPEN_TEXT,
    KIND_SINGLE_CHOICE,
    KIND_UNKNOWN,
    KNOWN_META_KEYS,
    SCHEMA_VERSION,
    new_source_metadata,
)

logger = logging.getLogger(__name__)


def _normalize_cell(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _classify_subheader(header: str, subheader: str) -> str:
    sub = subheader.strip()
    if header in KNOWN_META_KEYS or sub in KNOWN_META_KEYS:
        return KIND_META
    if sub == "Open-Ended Response":
        return KIND_OPEN_TEXT
    if sub == "Response":
        return KIND_SINGLE_CHOICE
    if sub:
        return KIND_MULTI_SELECT_OPTION
    return KIND_UNKNOWN


def _forward_fill_headers(headers: list[str]) -> list[str]:
    filled: list[str] = []
    last = ""
    for h in headers:
        h = h.strip()
        if h:
            last = h
        filled.append(last)
    return filled


def _build_groups(columns: list[ExportColumn]) -> list[ExportGroup]:
    groups: list[ExportGroup] = []
    by_header: dict[str, list[ExportColumn]] = {}
    for col in columns:
        if col.kind != KIND_MULTI_SELECT_OPTION:
            continue
        key = col.header.strip()
        if not key:
            continue
        by_header.setdefault(key, []).append(col)

    grp_idx = 0
    for header, cols in by_header.items():
        if len(cols) < 2:
            continue
        grp_idx += 1
        gid = f"grp_{grp_idx:03d}"
        col_ids = [c.id for c in cols]
        page_hint = cols[0].page_hint
        for c in cols:
            c.group_id = gid
        groups.append(
            ExportGroup(
                id=gid,
                header=header,
                kind="multi_select",
                column_ids=col_ids,
                page_hint=page_hint,
            )
        )
    return groups


def parse_excel_workbook(
    path: Path | str,
    *,
    version: int = 1,
    notes: str = "",
) -> ExportFormatDescriptor:
    """Parse row 1 (stems) and row 2 (subheaders) from the first worksheet."""
    try:
        import openpyxl
    except ImportError as e:
        raise RuntimeError(
            "openpyxl is not installed. Install it with:\n\n    pip install openpyxl\n"
        ) from e

    path = Path(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        row1 = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        row2 = next(ws.iter_rows(min_row=2, max_row=2, values_only=True), None)
        if not row1 or not row2:
            raise ValueError("Worksheet must have at least two header rows")

        raw_headers = [_normalize_cell(v) for v in row1]
        subheaders = [_normalize_cell(v) for v in row2]
        width = max(len(raw_headers), len(subheaders))
        raw_headers.extend([""] * (width - len(raw_headers)))
        subheaders.extend([""] * (width - len(subheaders)))

        filled_headers = _forward_fill_headers(raw_headers)
        warnings: list[str] = []
        if any(not h for h in raw_headers):
            warnings.append(
                "Sparse row-1 headers: question text appears only on first column of each block"
            )

        meta_columns: list[dict] = []
        columns: list[ExportColumn] = []
        col_idx = 0
        for ordinal, (header, sub) in enumerate(zip(filled_headers, subheaders), start=1):
            if not header and not sub:
                continue
            kind = _classify_subheader(header, sub)
            col_idx += 1
            col_id = f"col_{col_idx:03d}"
            if kind == KIND_META:
                meta_columns.append({"key": header or sub, "role": "system"})
            columns.append(
                ExportColumn(
                    id=col_id,
                    ordinal=ordinal,
                    header=header,
                    subheader=sub,
                    kind=kind,
                )
            )

        groups = _build_groups(columns)
        return ExportFormatDescriptor(
            schema_version=SCHEMA_VERSION,
            version=version,
            source=new_source_metadata(path.name, notes=notes),
            meta_columns=meta_columns,
            columns=columns,
            groups=groups,
            warnings_from_ingest=warnings,
        )
    finally:
        wb.close()
