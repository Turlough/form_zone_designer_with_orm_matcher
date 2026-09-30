"""Generate EXAMPLES/<survey>/export_config.json from json + Qualtrics template xlsx.

Usage (repo root, venv):

    python EXAMPLES/utils/generate_export_config.py EXAMPLES/RoI
    python EXAMPLES/utils/generate_export_config.py EXAMPLES/NI

Per-survey fixes the data cannot supply live in ``SURVEYS`` (keyed by survey folder name).
"""

from __future__ import annotations

import argparse
import difflib
import json
import re
import sys
from pathlib import Path

_REPO = Path(__file__).resolve().parents[2]
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))

from fields import RadioGroup, Tickbox
from util.designer_persistence import iter_runtime_fields
from util.field_metadata import export_display_title
from EXAMPLES.utils.qualtrics_headers import load_qualtrics_headers

ROI_COLUMNS: dict[int, str] = {
    10: "",  # rq_flag — blank insert
    14: "1.2 Age group",
    38: "Milking Platform - Owned acres",
    39: "Milking Platform - Leased acres (short-term)",
    40: "Milking Platform - Leased acres (medium-term)",
    41: "Milking Platform - Leased acres (long-term)",
    42: "Milking Platform - Total acres",
    43: "Outside Blocks- Owned acres",
    44: "Outside Blocks- Leased acres (short-term)",
    45: "Outside Blocks- Leased acres (medium-term)",
    46: "Outside Blocks- Leased acres (long-term)",
    47: "Outside Blocks- Total acres",
    48: "Total acres within 3km",
    49: "Only dairy cows on milking platform",
    55: "Q2_6_No_Issue_Not_Looking_To_Employ",
    60: "Q2_6_Challenges_Employing_International_Staff",
    62: "Q2_6_HR_Tax_And_Payroll_Implications",
    63: "Q2_6_Other_Comment",
    65: "Reason: Time for retirement",
    68: "Reason: Unable to meet investment",
    71: "Reason: Other (Details)",
    72: "Explore share farming options",
    74: "Use land for another farming enterprise",
    75: "Use land for alternative enterprise",
    78: "3.1 Outlook on future of milk production",
    80: "Cows milked in 2025",
    81: "Cows_2026",
    82: "Cows_2027",
    83: "Cows_2028",
    84: "Cows_2029",
    85: "Cows_2030",
    86: "Cows_2031",
    87: "TotalLitres_2026",
    88: "TotalLitres_2027",
    89: "TotalLitres_2028",
    90: "TotalLitres_2029",
    91: "TotalLitres_2030",
    92: "TotalLitres_2031",
    149: "In-calf heifers due to calve in 2026 (12+ months)",
    147: "Herd_Pct_Dairy_Breed_2026",
    98: "Factor: Limited farm infrastructure",
    99: "Factor: Employee availability",
    100: "Factor: Environmental issues and uncertainty",
    111: "Welcome conversation with Tirlán team",
    112: "What is your Whole Farm Stocking Rate (kg N/ha), as per the Nitrates Directive before any export of slurry for 2026? (e.g. Total whole farm kg N divided by number of hectares farmed)",
    116: "The Nitrates Derogation is in place until the end of 2028. With this extension, which best describes your approach?",
    117: "Action: Not applicable (Nitrates Derogation)",
    124: "Action: Other (comment)",
    125: "Are you currently exporting slurry?",
    127: "Will you practice zero grazing",
    128: "Dairy only (Cows, heifers, calves)",
    134: "Other",
    146: "",  # Dovea Genetics — online-only optional tick; no scan field
    188: "Are you aware of the new requirements",
    191: "6.14 None of these - Tick box",
    196: "6.14 Cow housing - All year round indoor system - Tick box",
    208: "6.14 Cow housing – All year round indoor system - Amount",
    218: "How likely are you to recommend Tirlan",
    221: "Welcome Tirlán conversation",
    225: "Signed",
    226: "County Letter",
    227: "6-Digit Herd Number",
    229: "Final Full Name",
}

SURVEYS: dict[str, dict] = {
    "RoI": {
        "description": "IFAC TirlanSurvey RoI 2026 — Qualtrics-compatible delivery transform",
        # template column -> field.name ("" = leave blank)
        "columns": ROI_COLUMNS,
        # Tickbox checked values the online data cannot supply (single scan tick vs online radio, consent).
        "tickbox_checked": {
            "Co_owner_non_family_Part_time": "Part time",  # no online responses in sample
            "Other": "Other",  # online column is free text; scan has only a tickbox
            "Is the land within 3km": "Yes",
            "Welcome conversation with Tirlán team": "Yes",
            "Welcome Tirlán conversation": "Yes",
            "Signed": "I consent",
        },
        # Radio labels whose online wording differs too much for fuzzy matching.
        "radio": {
            "Under the Department of Agriculture's Nitrates Banding System, which band do you fall into?": {
                "Q5.2 Band 1 (< 4,500kg milk)": "Band 1 (Less than 4,500kg milk)",
                "Q5.2 Band 2 (4,501-6,500kg milk)": "Band 2 (4,501kg – 6,500kg milk)",
                "Q5.2 Band 3 (> 6,500kg milk)": "Band 3 (Greater than 6,500kg milk)",
                "Q5.2 Don't know": "Don’t know",
            },
        },
        "merge_columns": [
            {
                "target_template_col": 223,
                "sources": ["Comments on any aspect of Tirlan", "Comments continued"],
                "separator": " ",
            }
        ],
    },
    "NI": {
        "description": "IFAC TirlanSurvey NI 2026 — Qualtrics-compatible delivery transform",
        # JSON field order matches the Excel survey columns once these print-only fields are skipped.
        "match": "order",
        "skip_fields": [
            "2025 Litres Supplied",
            "Challenge: Other factor selected",
            "Reasons - Other checkbox",
            "What is your farm's current status regarding the Nitrates Derogation?",
            "Looking ahead to the next 5 years, which best describes your position?",
            "No Dairy replacements on Farm",
        ],
        "columns": {10: ""},  # rq_flag
        # No online responses in the NI sample for these; values follow the RoI online convention.
        "tickbox_checked": {
            "Q1_4_Relief_Contract_workers_Full_time": "Full time",
            "Q1_4_Sharemilker_Full_time": "Full time",
            "Q1_4_Sharemilker_Part_time": "Part time",
            "Q1_4_Co_owner_non_family_Full_time": "Full time",
            "Q1_4_Co_owner_non_family_Part_time": "Part time",
            "None of the above - Have undertaken": "Have undertaken",
            "Other actions": "Other",  # online column is free text; scan has only a tickbox
            "Other": "Other",  # online column is free text; scan has only a tickbox
        },
        "radio": {
            "With respect to your non-dairy livestock, what are your plans over the next five years?": {
                "Non-dairy livestock plans - Reduce in numbers": "Reduce in numbers",
                "Non-dairy livestock plans - No livestock": "I do not have non-dairy livestock",
                "Non-dairy livestock plans - Increase in numbers": "Increase in numbers",
                "Non-dairy livestock plans - Remain at current levels": "Remain at current levels",
            },
            "6.8. In a normal year (i.e. normal weather, grass growth etc.), including dry cow feeding, how many Kg of concentrates/supplementary feed would you feed to each dairy cow?": {
                "Feed_Concentrates_Less_500kg_hd": "Less than 500kg/hd",  # RoI wording; not in NI sample
                "Feed_Concentrates_500_1000kg_hd": "500kg/hd to 1,000kg/hd",
                "Feed_Concentrates_1001_1500kg_hd": "1,001kg/hd to 1,500kg/hd",
                "Feed_Concentrates_1501_2000kg_hd": "1,501kg/hd to 2,000kg/hd",
                "Feed_Concentrates_2001_3000kg_hd": "2,001kg/hd to 3,000kg/hd",
                "Feed_Concentrates_3001_4000kg_hd": "3,001kg/hd to 4,000kg/hd",
                "Feed_Concentrates_More_4000kg_hd": "4,001kg/hd+",  # guessed from RoI "2,501kg/hd+"; not in NI sample
            },
        },
        "merge_columns": [],
    },
}

RADIO_MATCH_MIN_RATIO = 0.75
_LABEL_PREFIX = re.compile(r"^(?:q?\d+(?:\.\d+)*\.?\s+|[^:]{1,30}:\s*)", re.IGNORECASE)


def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", (s or "").strip().lower()).strip()


def _score_match(sub: str, stem: str, name: str, title: str) -> int:
    t, n, s, st = _norm(title), _norm(name), _norm(sub), _norm(stem)
    if s and (t == s or n == s):
        return 100
    if st and t == st:
        return 95
    if s and n.replace("_", " ") == s:
        return 90
    if s and all(w in t for w in s.split() if len(w) > 2):
        return 70
    return 0


def build_column_sources(
    json_dir: Path, xlsx: Path, manual_columns: dict[int, str], reserved: set[str] = frozenset()
) -> dict[str, str | None]:
    """Template column -> field.name. ``reserved`` fields (merge sources) are never auto-assigned."""
    headers = load_qualtrics_headers(xlsx)
    fields = [
        {
            "name": f.name.strip(),
            "title": export_display_title(f),
        }
        for _, f in iter_runtime_fields(json_dir)
    ]
    available = {f["name"] for f in fields} - set(reserved)

    assignments: dict[int, str | None] = {}
    for col_num, name in manual_columns.items():
        if name == "":
            assignments[col_num] = None
            continue
        if name in available:
            assignments[col_num] = name
            available.discard(name)

    for col in headers.columns:
        if col.col in assignments:
            continue
        candidates: list[tuple[int, str]] = []
        for f in fields:
            if f["name"] not in available:
                continue
            sc = _score_match(col.sub, col.stem, f["name"], f["title"])
            if sc >= 70:
                candidates.append((sc, f["name"]))
        candidates.sort(reverse=True)
        if candidates:
            pick = candidates[0][1]
            assignments[col.col] = pick
            available.discard(pick)
        else:
            assignments[col.col] = None

    return {str(k): v for k, v in sorted(assignments.items())}


def build_column_sources_by_order(
    json_dir: Path, xlsx: Path, manual_columns: dict[int, str], skip_fields: list[str]
) -> dict[str, str | None]:
    """Pair template columns with JSON fields in order (manual columns and skip_fields excluded)."""
    headers = load_qualtrics_headers(xlsx)
    skip = set(skip_fields) | {n for n in manual_columns.values() if n}
    names = [f.name.strip() for _, f in iter_runtime_fields(json_dir) if f.name.strip() not in skip]
    cols = [c.col for c in headers.columns if c.col not in manual_columns]
    if len(names) != len(cols):
        raise SystemExit(
            f"Order match needs equal counts: {len(cols)} template columns vs {len(names)} fields. "
            "Adjust skip_fields."
        )
    assignments: dict[int, str | None] = {c: (n or None) for c, n in manual_columns.items()}
    assignments.update(zip(cols, names))
    return {str(k): v for k, v in sorted(assignments.items())}


def _online_values(xlsx: Path) -> dict[int, set[str]]:
    """Distinct non-empty values per 1-based column in the template's data rows (row 3+)."""
    from openpyxl import load_workbook

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        values: dict[int, set[str]] = {}
        for row in ws.iter_rows(min_row=3, values_only=True):
            for i, v in enumerate(row, start=1):
                if v is not None and str(v).strip():
                    values.setdefault(i, set()).add(str(v).strip())
    finally:
        wb.close()
    return values


def _label_key(text: str) -> str:
    text = _LABEL_PREFIX.sub("", (text or "").strip())
    text = text.replace("’", "'").replace("‘", "'").replace("–", "-").replace("—", "-")
    text = re.sub(r"\s*([-/])\s*", r"\1", text)
    return re.sub(r"\s+", " ", text).casefold()


def _best_online_label(label: str, candidates: set[str]) -> tuple[str | None, float]:
    key = _label_key(label)
    best, best_ratio = None, 0.0
    for cand in candidates:
        ratio = difflib.SequenceMatcher(None, key, _label_key(cand)).ratio()
        if ratio > best_ratio:
            best, best_ratio = cand, ratio
    return best, best_ratio


def build_value_maps(
    json_dir: Path,
    xlsx: Path,
    column_sources: dict[str, str | None],
    manual_tickbox: dict[str, str],
    manual_radio: dict[str, dict[str, str]],
) -> tuple[dict[str, str], dict[str, dict[str, str]], list[str]]:
    """Tickbox checked values and radio label maps that make scan output match online rows."""
    subs = {c.col: c.sub for c in load_qualtrics_headers(xlsx).columns}
    online = _online_values(xlsx)
    fields = {f.name.strip(): f for _, f in iter_runtime_fields(json_dir)}

    tickbox_checked: dict[str, str] = {}
    radio: dict[str, dict[str, str]] = {}
    warnings: list[str] = []

    for key, name in column_sources.items():
        field = fields.get(name or "")
        if field is None:
            continue
        col = int(key)
        vals = online.get(col, set())

        if isinstance(field, Tickbox):
            if name in manual_tickbox:
                tickbox_checked[name] = manual_tickbox[name]
            elif len(vals) == 1 and next(iter(vals)) != subs.get(col, ""):
                tickbox_checked[name] = next(iter(vals))
            continue

        if not isinstance(field, RadioGroup) or not vals:
            continue
        manual = manual_radio.get(name, {})
        labels = [(b.name or "").strip() for b in field.radio_buttons]
        # An online label equal to one scan label belongs to that button; never fuzzy-match onto it.
        unclaimed = vals - set(labels) - set(manual.values())
        mapping: dict[str, str] = {}
        for label in labels:
            if not label or label in vals:
                continue
            if label in manual:
                mapping[label] = manual[label]
                continue
            best, ratio = _best_online_label(label, unclaimed)
            if best is not None and ratio >= RADIO_MATCH_MIN_RATIO:
                mapping[label] = best
            else:
                warnings.append(f"c{col} {name!r}: no online match for {label!r}")
        if len(set(mapping.values())) != len(mapping):
            warnings.append(f"c{col} {name!r}: several scan labels map to one online label")
        if mapping:
            radio[name] = mapping

    return tickbox_checked, radio, warnings


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("survey_dir", type=Path)
    args = parser.parse_args(argv)
    survey_dir = args.survey_dir.resolve()
    json_dir = survey_dir / "json"
    xlsx_files = sorted(survey_dir.glob("*.xlsx"))
    if not json_dir.is_dir() or not xlsx_files:
        raise SystemExit(f"Need json/ and a .xlsx in {survey_dir}")

    survey = SURVEYS.get(survey_dir.name)
    if survey is None:
        raise SystemExit(f"Add a SURVEYS entry for {survey_dir.name!r} in {Path(__file__).name}")

    merge_columns = survey["merge_columns"]
    merge_sources = {s for spec in merge_columns for s in spec["sources"]}
    template_xlsx = xlsx_files[0].name
    if survey.get("match") == "order":
        column_sources = build_column_sources_by_order(
            json_dir, xlsx_files[0], survey["columns"], survey.get("skip_fields", [])
        )
    else:
        column_sources = build_column_sources(json_dir, xlsx_files[0], survey["columns"], merge_sources)
    for spec in merge_columns:
        column_sources[str(spec["target_template_col"])] = None
    tickbox_checked, radio_maps, warnings = build_value_maps(
        json_dir, xlsx_files[0], column_sources, survey["tickbox_checked"], survey["radio"]
    )

    used = {v for v in column_sources.values() if v}
    all_names = {f.name.strip() for _, f in iter_runtime_fields(json_dir)}
    remove_fields = sorted(n for n in all_names if n and n not in used and n not in merge_sources)

    config = {
        "format_version": 1,
        "description": survey["description"],
        "template_xlsx": template_xlsx,
        "json_folder": "json",
        "first_column": "File",
        "prepend_qualtrics_metadata": True,
        "output_suffix": "_for_import",
        "tickbox_default_checked": "use_template_subheader",
        "value_maps": {
            "tickbox_checked": tickbox_checked,
            "radio": radio_maps,
        },
        "merge_columns": merge_columns,
        "remove_fields": remove_fields,
        "column_sources": column_sources,
    }

    out_path = survey_dir / "export_config.json"
    out_path.write_text(json.dumps(config, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Wrote {out_path} ({len(column_sources)} template columns, {len(remove_fields)} remove_fields)")
    unmapped = [k for k, v in column_sources.items() if v is None]
    if unmapped:
        print(f"Blank template columns ({len(unmapped)}): {unmapped[:12]}{'…' if len(unmapped) > 12 else ''}")
    print(f"value_maps: {len(tickbox_checked)} tickbox_checked, {len(radio_maps)} radio fields")
    for w in warnings:
        print(f"WARNING {w}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
