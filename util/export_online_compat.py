"""Transform Exporter delivery CSV into Qualtrics-shaped Excel using export_config.json."""

from __future__ import annotations

import csv
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from fields import Tickbox, SignatureField
from util.designer_persistence import export_title_map, iter_runtime_fields

TICKBOX_TRUTHY = frozenset(
    {"ticked", "signed", "true", "1", "yes", "checked", "tick", "y"}
)


def _norm_key(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "").strip()).casefold()


def load_export_config(path: Path) -> dict[str, Any]:
    with path.open(encoding="utf-8") as f:
        data = json.load(f)
    if data.get("format_version") != 1:
        raise ValueError(f"Unsupported export_config format_version: {data.get('format_version')!r}")
    return data


def resolve_template_path(
    config_path: Path,
    config: dict[str, Any],
    *,
    template_override: Path | None = None,
) -> Path:
    """Resolve Qualtrics template workbook (dialog override or export_config relative path)."""
    if template_override is not None:
        template_path = template_override.resolve()
        if not template_path.is_file():
            raise FileNotFoundError(f"Template workbook not found: {template_path}")
        return template_path

    base = config_path.parent
    template_name = (config.get("template_xlsx") or "").strip()
    if template_name:
        template_path = (base / template_name).resolve()
        if template_path.is_file():
            return template_path
    raise FileNotFoundError(
        f"No Excel template found beside {config_path.name}. "
        "Choose the Qualtrics export .xlsx in the dialog (expected in the config folder)."
    )


def resolve_config_path(
    config_path: Path,
    *,
    template_override: Path | None = None,
) -> tuple[Path, dict[str, Any]]:
    config_path = config_path.resolve()
    config = load_export_config(config_path)
    template_path = resolve_template_path(config_path, config, template_override=template_override)
    return template_path, config


def _load_qualtrics_template_layout(
    template_path: Path, *, sheet: str | None = None
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Return (prefix meta + rq_flag, survey columns) from the Qualtrics template."""
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise RuntimeError(
            "openpyxl is required for online export compatibility. Install openpyxl in the project venv."
        ) from e

    wb = load_workbook(template_path, read_only=True, data_only=True)
    try:
        name = sheet or wb.sheetnames[0]
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    if len(rows) < 2:
        raise ValueError(f"{template_path} must have at least two header rows")

    def cell(row: list, i: int) -> str:
        if i >= len(row) or row[i] is None:
            return ""
        return str(row[i]).strip()

    row1 = [cell(rows[0], i) for i in range(max(len(rows[0]), len(rows[1])))]
    row2 = [cell(rows[1], i) for i in range(len(row1))]
    while len(row2) < len(row1):
        row2.append("")

    filled: list[str] = []
    last = ""
    for h in row1:
        if h:
            last = h
        filled.append(last)

    meta_stems = frozenset(
        {
            "Respondent ID",
            "Collector ID",
            "Start Date",
            "End Date",
            "IP Address",
            "Email Address",
            "First Name",
            "Last Name",
            "Custom Data 1",
        }
    )

    def _kind(sub: str) -> str:
        if sub == "Response":
            return "radio"
        if sub == "Open-Ended Response":
            return "text"
        if sub:
            return "named"
        return "empty"

    prefix: list[dict[str, Any]] = []
    survey: list[dict[str, Any]] = []
    for i, (stem, sub) in enumerate(zip(filled, row2), start=1):
        if not stem and not sub:
            continue
        entry = {
            "template_col": i,
            "stem": stem,
            "sub": sub,
            "kind": _kind(sub),
        }
        if stem in meta_stems or stem.casefold() == "rq_flag":
            prefix.append(entry)
        else:
            survey.append(entry)
    return prefix, survey


def _load_qualtrics_template_columns(template_path: Path, *, sheet: str | None = None) -> list[dict[str, Any]]:
    """Survey columns only (legacy helper)."""
    _, survey = _load_qualtrics_template_layout(template_path, sheet=sheet)
    return survey


@dataclass(frozen=True)
class _FieldInfo:
    name: str
    delivery_title: str
    is_tick: bool
    is_signature: bool


def _field_infos(json_folder: Path) -> dict[str, _FieldInfo]:
    title_map = export_title_map(json_folder)
    infos: dict[str, _FieldInfo] = {}
    for _, field in iter_runtime_fields(json_folder):
        name = (getattr(field, "name", None) or "").strip()
        if not name:
            continue
        delivery = title_map.get(name, name)
        is_tick = isinstance(field, Tickbox) and not isinstance(field, SignatureField)
        is_sig = isinstance(field, SignatureField)
        infos[name] = _FieldInfo(name=name, delivery_title=delivery, is_tick=is_tick, is_signature=is_sig)
    return infos


def _is_truthy_tick(raw: str, checked_value: str = "Ticked") -> bool:
    text = (raw or "").strip()
    if not text:
        return False
    low = text.casefold()
    if low in TICKBOX_TRUTHY:
        return True
    if checked_value and low == checked_value.casefold():
        return True
    return False


def _load_value_maps(config: dict[str, Any]) -> tuple[dict[str, str], dict[str, dict[str, str]]]:
    """``value_maps.tickbox_checked`` (field → online value) and ``value_maps.radio``
    (field → scan label → online label). Radio keys are matched case-insensitively."""
    section = config.get("value_maps") or {}
    tickbox_checked = {str(k): str(v) for k, v in (section.get("tickbox_checked") or {}).items()}
    radio: dict[str, dict[str, str]] = {}
    for field_name, mapping in (section.get("radio") or {}).items():
        if isinstance(mapping, dict):
            radio[str(field_name)] = {str(a).strip().casefold(): str(b) for a, b in mapping.items()}
    return tickbox_checked, radio


def _transform_cell(
    raw: str,
    *,
    field_name: str | None,
    template_sub: str,
    field_info: _FieldInfo | None,
    tickbox_checked: dict[str, str],
    tickbox_default: str,
    radio_maps: dict[str, dict[str, str]],
) -> str:
    text = (raw or "").strip()

    if field_info and (field_info.is_tick or field_info.is_signature):
        if not _is_truthy_tick(text):
            return ""
        if field_name and field_name in tickbox_checked:
            return tickbox_checked[field_name]
        if tickbox_default == "use_template_subheader" and template_sub:
            return template_sub
        return text or "Ticked"

    mapping = radio_maps.get(field_name or "")
    if mapping and text:
        return mapping.get(text.casefold(), text)
    return text


def _read_csv_rows(csv_path: Path) -> tuple[list[str], list[list[str]]]:
    with csv_path.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        headers = next(reader, None)
        if not headers:
            raise ValueError(f"CSV has no header row: {csv_path}")
        rows = [row for row in reader]
    return headers, rows


def _header_index(headers: list[str]) -> dict[str, int]:
    return {_norm_key(h): i for i, h in enumerate(headers)}


def _merge_comment_values(
    row: list[str],
    header_idx: dict[str, int],
    merge_spec: dict[str, Any],
    field_infos: dict[str, _FieldInfo],
) -> str:
    parts: list[str] = []
    for source in merge_spec.get("sources") or []:
        source = str(source).strip()
        if not source:
            continue
        idx = header_idx.get(_norm_key(source))
        if idx is None and source in field_infos:
            idx = header_idx.get(_norm_key(field_infos[source].delivery_title))
        if idx is None or idx >= len(row):
            continue
        text = (row[idx] or "").strip()
        if text:
            parts.append(text)
    sep = merge_spec.get("separator", " ")
    return sep.join(parts)


def transform_delivery_csv_to_xlsx(
    *,
    csv_path: Path,
    config_path: Path,
    template_path: Path | None = None,
    output_path: Path | None = None,
) -> Path:
    """Write Qualtrics-layout xlsx from a delivery CSV and export_config.json."""
    try:
        from openpyxl import Workbook
    except ImportError as e:
        raise RuntimeError(
            "openpyxl is required for online export compatibility. Install openpyxl in the project venv."
        ) from e

    csv_path = csv_path.resolve()
    config_path = config_path.resolve()
    template_resolved, config = resolve_config_path(
        config_path,
        template_override=template_path,
    )
    base = config_path.parent

    json_rel = (config.get("json_folder") or "json").strip()
    json_folder = (base / json_rel).resolve()
    if not json_folder.is_dir():
        raise FileNotFoundError(f"json_folder not found: {json_folder}")

    prefix_columns, template_columns = _load_qualtrics_template_layout(
        template_resolved, sheet=config.get("template_sheet")
    )
    prepend_meta = config.get("prepend_qualtrics_metadata", True)
    if not prepend_meta:
        prefix_columns = []
    column_sources: dict[str, str | None] = {
        str(k): v for k, v in (config.get("column_sources") or {}).items()
    }

    field_infos = _field_infos(json_folder)
    tickbox_checked, radio_maps = _load_value_maps(config)
    tickbox_default = (config.get("tickbox_default_checked") or "use_template_subheader").strip()

    merge_specs = config.get("merge_columns") or []
    first_column = (config.get("first_column") or "File").strip()

    headers, data_rows = _read_csv_rows(csv_path)
    header_idx = _header_index(headers)
    file_idx = header_idx.get(_norm_key(first_column))
    if file_idx is None:
        for alt in ("file", "tiff_path", "path", "document_path"):
            file_idx = header_idx.get(alt)
            if file_idx is not None:
                break
    if file_idx is None:
        raise ValueError(f"CSV missing first column {first_column!r}")

    merge_by_target: dict[int, dict[str, Any]] = {}
    for spec in merge_specs:
        target = spec.get("target_template_col")
        if target is not None:
            merge_by_target[int(target)] = spec

    out_path = output_path
    if out_path is None:
        suffix = config.get("output_suffix") or "_for_import"
        out_path = csv_path.with_name(f"{csv_path.stem}{suffix}.xlsx")
    out_path = out_path.resolve()

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    out_headers_row1: list[str] = [first_column]
    out_headers_row2: list[str] = [""]
    for col in prefix_columns:
        out_headers_row1.append(col["stem"])
        out_headers_row2.append(col["sub"])
    for col in template_columns:
        out_headers_row1.append(col["stem"])
        out_headers_row2.append(col["sub"])

    ws.append(out_headers_row1)
    ws.append(out_headers_row2)

    for row in data_rows:
        out_row: list[str] = []
        file_val = row[file_idx] if file_idx < len(row) else ""
        out_row.append(file_val.strip())
        out_row.extend([""] * len(prefix_columns))

        for col in template_columns:
            tcol = col["template_col"]
            key = str(tcol)
            source_name = column_sources.get(key)
            if tcol in merge_by_target:
                out_row.append(
                    _merge_comment_values(row, header_idx, merge_by_target[tcol], field_infos)
                )
                continue
            if not source_name:
                out_row.append("")
                continue

            info = field_infos.get(source_name)
            delivery_title = info.delivery_title if info else source_name
            idx = header_idx.get(_norm_key(delivery_title))
            if idx is None:
                idx = header_idx.get(_norm_key(source_name))
            raw = row[idx] if idx is not None and idx < len(row) else ""
            out_row.append(
                _transform_cell(
                    raw,
                    field_name=source_name,
                    template_sub=col.get("sub") or "",
                    field_info=info,
                    tickbox_checked=tickbox_checked,
                    tickbox_default=tickbox_default,
                    radio_maps=radio_maps,
                )
            )

        ws.append(out_row)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def default_output_path_for_csv(csv_path: Path, config_path: Path) -> Path:
    config = load_export_config(config_path.resolve())
    suffix = config.get("output_suffix") or "_for_import"
    csv_path = csv_path.resolve()
    return csv_path.with_name(f"{csv_path.stem}{suffix}.xlsx")
